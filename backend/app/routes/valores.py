"""Gestão de valores (painel estilo Power BI) — agregações financeiras."""
from fastapi import APIRouter, Depends
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.middleware import get_current_user
from app.models import User, Processo, NotaFiscal, LancamentoBancario, Coletador
from app.services import get_db

router = APIRouter(prefix="/api/v1/valores", tags=["valores"])


def exigir_gestao(user: User = Depends(get_current_user)) -> User:
    from fastapi import HTTPException
    if user.role.name not in ("admin", "power_user"):
        raise HTTPException(status_code=403, detail="Sem permissão")
    return user


@router.get("/resumo")
async def resumo(db: Session = Depends(get_db), user: User = Depends(exigir_gestao)):
    # honorários por status
    por_status = (
        db.query(Processo.status, func.count(Processo.id), func.coalesce(func.sum(Processo.honorarios), 0))
        .group_by(Processo.status).all()
    )
    # honorários por SETOR (limpo) — antes agrupava por especialidade (texto podre)
    por_area = (
        db.query(Processo.setor, func.coalesce(func.sum(Processo.honorarios), 0))
        .filter(Processo.honorarios.isnot(None))
        .group_by(Processo.setor).all()
    )
    # notas fiscais por status
    notas = db.query(NotaFiscal.status, func.count(NotaFiscal.id),
                     func.coalesce(func.sum(NotaFiscal.valor_servico), 0)).group_by(NotaFiscal.status).all()
    # conciliação: conciliado x pendente
    conc = db.query(LancamentoBancario.status, func.count(LancamentoBancario.id),
                    func.coalesce(func.sum(LancamentoBancario.valor), 0)).group_by(LancamentoBancario.status).all()

    # evolução mensal de honorários (por mês de criação do processo)
    from sqlalchemy import extract
    por_mes_raw = (
        db.query(extract("year", Processo.created_at), extract("month", Processo.created_at),
                 func.coalesce(func.sum(Processo.honorarios), 0), func.count(Processo.id))
        .filter(Processo.honorarios.isnot(None))
        .group_by(extract("year", Processo.created_at), extract("month", Processo.created_at))
        .order_by(extract("year", Processo.created_at), extract("month", Processo.created_at)).all()
    )
    por_mes = [{"label": f"{int(m):02d}/{int(a)}", "valor": float(v), "qtd": int(q)}
               for a, m, v, q in por_mes_raw if a]

    total_honorarios = float(db.query(func.coalesce(func.sum(Processo.honorarios), 0)).scalar() or 0)
    total_notas = float(db.query(func.coalesce(func.sum(NotaFiscal.valor_servico), 0))
                        .filter(NotaFiscal.status == "emitida").scalar() or 0)
    total_coletadores = db.query(func.count(Coletador.id)).scalar() or 0
    total_processos = db.query(func.count(Processo.id)).scalar() or 0

    return {
        "cartoes": {
            "total_honorarios": total_honorarios,
            "total_notas_emitidas": total_notas,
            "processos": total_processos,
            "coletadores": total_coletadores,
        },
        "honorarios_por_mes": por_mes,
        "honorarios_por_status": [
            {"label": s or "Sem status", "qtd": int(q), "valor": float(v)} for s, q, v in por_status
        ],
        "honorarios_por_area": [
            {"label": a or "Sem classificação", "valor": float(v)} for a, v in por_area
        ],
        "notas": [
            {"label": s, "qtd": int(q), "valor": float(v)} for s, q, v in notas
        ],
        "conciliacao": [
            {"label": s, "qtd": int(q), "valor": float(v)} for s, q, v in conc
        ],
    }


def _dados_relatorio(db, tipo: str):
    """Retorna (colunas, linhas) do relatório escolhido."""
    if tipo == "processos":
        rows = db.query(Processo).all()
        cols = ["Nº Processo", "Área", "Status", "Responsável", "Honorários", "Pago"]
        linhas = [[p.numero_cnj, p.especialidade or "", p.status or "", p.responsavel or "",
                   float(p.honorarios or 0), "Sim" if p.pago else "Não"] for p in rows]
    elif tipo == "notas":
        rows = db.query(NotaFiscal).all()
        cols = ["ID", "Processo", "Status", "Nº NFSe", "Valor"]
        linhas = [[n.id, n.processo_id, n.status, n.numero_nfse or "", float(n.valor_servico or 0)] for n in rows]
    elif tipo == "coletadores":
        rows = db.query(Coletador).all()
        cols = ["Código SCPG", "Nome", "CPF", "Ativo"]
        linhas = [[c.codigo_scpg, c.nome_completo, c.cpf, "Sim" if c.ativo else "Não"] for c in rows]
    else:  # conciliacao
        rows = db.query(LancamentoBancario).all()
        cols = ["Data", "Favorecido", "Valor", "Status"]
        linhas = [[str(l.data or ""), (l.favorecido or "")[:40], float(l.valor or 0), l.status] for l in rows]
    return cols, linhas


@router.get("/exportar")
async def exportar(tipo: str = "processos", formato: str = "xlsx",
                   db: Session = Depends(get_db), user: User = Depends(exigir_gestao)):
    """Exporta um relatório em .xlsx ou .pdf. tipo: processos|notas|coletadores|conciliacao."""
    from fastapi.responses import StreamingResponse
    import io
    cols, linhas = _dados_relatorio(db, tipo)

    if formato == "pdf":
        from fpdf import FPDF
        import unicodedata
        def ascii_(s): return unicodedata.normalize("NFKD", str(s)).encode("latin-1", "ignore").decode("latin-1")
        pdf = FPDF(orientation="L")
        pdf.add_page(); pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 10, ascii_(f"Relatorio: {tipo}"), ln=True)
        pdf.set_font("Helvetica", "B", 8)
        larg = (pdf.w - 20) / len(cols)
        for c in cols: pdf.cell(larg, 7, ascii_(c)[:25], border=1)
        pdf.ln(); pdf.set_font("Helvetica", "", 8)
        for linha in linhas[:500]:
            for v in linha: pdf.cell(larg, 6, ascii_(v)[:25], border=1)
            pdf.ln()
        buf = io.BytesIO(pdf.output()); buf.seek(0)
        return StreamingResponse(buf, media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="relatorio_{tipo}.pdf"'})

    # xlsx
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = tipo[:31]
    ws.append(cols)
    for linha in linhas:
        ws.append(linha)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="relatorio_{tipo}.xlsx"'})
