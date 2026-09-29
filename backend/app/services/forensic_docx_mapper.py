"""
Mapper entre dados de análise forense e placeholders do template DOCX.
Extrai {{VARIAVEL}} e mapeia dados reais.
"""
from pathlib import Path
from docx import Document
import re
from typing import Dict, Set


class ForensicDocxMapper:
    """Extrai e mapeia placeholders do template DOCX de laudo forense."""

    TEMPLATE_PATH = Path("/var/www/perito-v6/backend/app/templates/TEMPLATE_LAUDO_FORENSE.docx")

    def __init__(self, template_path: Path = None):
        self.template_path = template_path or self.TEMPLATE_PATH
        self.doc = Document(self.template_path)
        self.placeholders: Set[str] = set()
        self._extract_all_placeholders()

    def _extract_all_placeholders(self) -> None:
        """Extrai todos os {{VARIAVEL}} do documento."""
        for para in self.doc.paragraphs:
            matches = re.findall(r'\{\{[A-Z_0-9]+\}\}', para.text)
            self.placeholders.update(matches)

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        matches = re.findall(r'\{\{[A-Z_0-9]+\}\}', para.text)
                        self.placeholders.update(matches)

    def extract_placeholders(self) -> Set[str]:
        """Retorna set de todos os placeholders encontrados."""
        return self.placeholders

    def map_data(self, data: Dict[str, str]) -> Dict[str, str]:
        """
        Mapeia dados para placeholders.

        Args:
            data: Dict com chaves minúsculas (e.g., 'codigo_laudo')

        Returns:
            Dict mapeado com chaves {{MAIUSCULA}} (e.g., {{CODIGO_LAUDO}})
        """
        mapping = {}
        for placeholder in self.placeholders:
            # Remove {{ e }} e converte para minúsculas
            key = placeholder.strip('{}').lower()
            if key in data:
                mapping[placeholder] = data[key]
            else:
                mapping[placeholder] = "[NÃO PREENCHIDO]"
        return mapping
