"""
Script para criar templates de ofício em Word (.docx) com placeholders {{CAMPO}}.
Roda uma vez na inicialização ou sob demanda.
"""
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def criar_templates():
    """Cria templates padrão de ofícios em DOCX."""
    templates_dir = os.path.join(os.path.dirname(__file__), "..", "templates")
    os.makedirs(templates_dir, exist_ok=True)

    # Template: Ofício de Requerimento
    criar_oficio_requerimento(os.path.join(templates_dir, "oficio_requerimento.docx"))

    # Template: Ofício de Manifestação
    criar_oficio_manifestacao(os.path.join(templates_dir, "oficio_manifestacao.docx"))

    # Template: Ofício Resposta Quesito
    criar_oficio_resposta_quesito(os.path.join(templates_dir, "oficio_resposta_quesito.docx"))


def criar_oficio_requerimento(caminho: str):
    """Template para requerimento de perícia."""
    doc = Document()

    # Cabeçalho
    doc.add_paragraph("MINISTÉRIO PÚBLICO DO ESTADO DE MATO GROSSO DO SUL", style="Heading 1").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("{{VARA}}", style="Heading 2").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    # Data e local
    doc.add_paragraph("{{DATA}}")
    doc.add_paragraph("")

    # Destinatário
    doc.add_paragraph("Excelentíssimo Senhor Juiz {{JUIZ}}")
    doc.add_paragraph("da {{VARA}} de {{COMARCA}}")
    doc.add_paragraph("Processo nº {{NUMERO}}")
    doc.add_paragraph("")

    # Assunto
    doc.add_paragraph("ASSUNTO: Requerimento de Realização de Perícia").bold = True
    doc.add_paragraph("")

    # Corpo
    corpo = doc.add_paragraph(
        """Pela presente, o Ministério Público vem requerer a Vossa Excelência a realização de perícia contábil/técnica no presente feito, a fim de esclarecer os pontos controversos entre as partes.

Objeto da perícia: {{RESUMO_INTIMACAO}}

Prazo requerido: {{PRAZO_DIAS}} dias corridos

Fundamentação: Art. 473, CPC e Lei nº 13.105/2015.

Nestes termos, pede deferimento.

Respeitosamente,"""
    )

    # Assinatura
    doc.add_paragraph("")
    doc.add_paragraph("_____________________________")
    doc.add_paragraph("{{RESPONSAVEL}}")
    doc.add_paragraph("OAB/MS nº")

    doc.save(caminho)
    print(f"✅ Template criado: {caminho}")


def criar_oficio_manifestacao(caminho: str):
    """Template para manifestação."""
    doc = Document()

    doc.add_paragraph("MINISTÉRIO PÚBLICO DO ESTADO DE MATO GROSSO DO SUL", style="Heading 1").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("{{VARA}}", style="Heading 2").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    doc.add_paragraph("{{DATA}}")
    doc.add_paragraph("")

    doc.add_paragraph("Excelentíssimo Senhor Juiz {{JUIZ}}")
    doc.add_paragraph("Processo nº {{NUMERO}}")
    doc.add_paragraph("")

    doc.add_paragraph("ASSUNTO: Manifestação sobre os autos").bold = True
    doc.add_paragraph("")

    doc.add_paragraph(
        """Manifestamos-nos nos autos em tela, como segue:

{{RESUMO_INTIMACAO}}

Fundamentação: CPC e leis aplicáveis.

Nestes termos, pede-se nota.

Respeitosamente,"""
    )

    doc.add_paragraph("")
    doc.add_paragraph("_____________________________")
    doc.add_paragraph("{{RESPONSAVEL}}")

    doc.save(caminho)
    print(f"✅ Template criado: {caminho}")


def criar_oficio_resposta_quesito(caminho: str):
    """Template para resposta a quesito."""
    doc = Document()

    doc.add_paragraph("RESPOSTA AOS QUESITOS", style="Heading 1").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    doc.add_paragraph("Processo: {{NUMERO}}")
    doc.add_paragraph("Juíz: {{JUIZ}}")
    doc.add_paragraph("Data: {{DATA_HOJE}}")
    doc.add_paragraph("")

    doc.add_paragraph("QUESITO 1:")
    doc.add_paragraph("Resposta: ")
    doc.add_paragraph("")

    doc.add_paragraph("QUESITO 2:")
    doc.add_paragraph("Resposta: ")
    doc.add_paragraph("")

    doc.add_paragraph("_____________________________")
    doc.add_paragraph("{{RESPONSAVEL}}")

    doc.save(caminho)
    print(f"✅ Template criado: {caminho}")


if __name__ == "__main__":
    criar_templates()
